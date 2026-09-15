import csv
import json
import shutil
import sys
from pathlib import Path
from unittest.mock import patch

import openpyxl

PROJECT_DIR = Path(__file__).resolve().parents[1]
AUTOMATION_DIR = PROJECT_DIR / "automation"
sys.path.insert(0, str(AUTOMATION_DIR))

import dashboard_integration as integration
import run_full_pipeline
from validate_dashboard_data import validate_dashboard_snapshot


def make_xlsx_report(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append(["MAKER", "JAN", "FEB"])
    for row in rows:
        ws.append(row)
    wb.save(path)


def test_normalize_rto_code_handles_spacing_and_zero_padding():
    assert integration.normalize_rto_code("KL 01") == "KL1"
    assert integration.normalize_rto_code("KL07") == "KL7"
    assert integration.normalize_rto_code("TRIVANDRUM - KL 1 (2026)") == "KL1"


def test_extract_report_rows_reads_html_table_saved_as_xlsx(tmp_path):
    report_path = tmp_path / "KL2.xlsx"
    report_path.write_text(
        """
        <table>
          <tr><th>Maker</th><th>JAN</th><th>FEB</th></tr>
          <tr><td>TOYOTA KIRLOSKAR MOTOR PVT LTD</td><td>10</td><td>12</td></tr>
          <tr><td>TOTAL</td><td>10</td><td>12</td></tr>
        </table>
        """,
        encoding="utf-8",
    )

    rows = integration.extract_report_rows(report_path, "KL2", "KOLLAM RTO")

    assert rows == [{
        "RTO": "KL2",
        "NAME": "KOLLAM RTO",
        "MAKER": "TOYOTA KIRLOSKAR MOTOR PVT LTD",
        "JAN": 10,
        "FEB": 12,
    }]


def test_extract_report_rows_preserves_zero_volume_maker_rows(tmp_path):
    report_path = tmp_path / "KL3.xlsx"
    make_xlsx_report(report_path, [["ZERO MAKER LTD", 0, 0]])

    rows = integration.extract_report_rows(report_path, "KL3", "PATHANAMTHITTA RTO")

    assert rows == [{
        "RTO": "KL3",
        "NAME": "PATHANAMTHITTA RTO",
        "MAKER": "ZERO MAKER LTD",
        "JAN": 0,
        "FEB": 0,
    }]


def test_extract_report_rows_reads_actual_vahan_multirow_header_layout(tmp_path):
    report_path = tmp_path / "KL1.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "reportTable"
    ws.append(["Maker Month Wise Data  of TRIVANDRUM RTO - KL1 , Kerala (2026)"])
    ws.append(["S No", " Maker ", "Month Wise ", None, None, None])
    ws.append([None, None, None, None, None, None])
    ws.append(["", "", "JAN", "FEB", "MAR", "APR"])
    ws.append(["1", "TOYOTA KIRLOSKAR MOTOR PVT LTD", "10", "11", "12", "13"])
    ws.append(["2", "ZERO MAKER LTD", "0", "0", "0", "0"])
    wb.save(report_path)

    rows = integration.extract_report_rows(report_path, "KL1", "TRIVANDRUM RTO")

    assert rows[0]["MAKER"] == "TOYOTA KIRLOSKAR MOTOR PVT LTD"
    assert rows[0]["JAN"] == 10
    assert rows[0]["APR"] == 13
    assert rows[1]["MAKER"] == "ZERO MAKER LTD"


def test_load_reports_fails_when_folder_has_no_reports(tmp_path):
    try:
        integration.load_reports(tmp_path, {})
    except FileNotFoundError as exc:
        assert "No report files found" in str(exc)
    else:
        raise AssertionError("Expected FileNotFoundError")


def test_load_reports_skips_out_of_scope_files_and_unapproved_makers(tmp_path):
    make_xlsx_report(tmp_path / "KL1.xlsx", [
        ["TOYOTA KIRLOSKAR MOTOR PVT LTD", 10, 20],
        ["TWO WHEELER MAKER LTD", 99, 99],
    ])
    make_xlsx_report(tmp_path / "KL99.xlsx", [
        ["TOYOTA KIRLOSKAR MOTOR PVT LTD", 1, 1],
    ])

    rows, files = integration.load_reports(tmp_path, {"KL1": "TRIVANDRUM RTO"})

    assert [file.name for file in files] == ["KL1.xlsx"]
    assert len(rows) == 1
    assert rows[0]["MAKER"] == "TOYOTA KIRLOSKAR MOTOR PVT LTD"
    assert rows[0]["RTO"] == "KL1"


def test_integrate_fails_before_writing_when_an_rto_is_missing(tmp_path):
    reports = tmp_path / "reports" / "2026-09-14"
    reports.mkdir(parents=True)
    for number in range(1, 86):
        make_xlsx_report(
            reports / f"KL{number}.xlsx",
            [["TOYOTA KIRLOSKAR MOTOR PVT LTD", 10, 20]],
        )

    with patch.object(integration, "copy_raw_reports") as copy_raw_reports:
        try:
            integration.integrate(reports_dir=reports, dashboard_dir=tmp_path)
        except RuntimeError as exc:
            assert "Expected 86" in str(exc)
        else:
            raise AssertionError("Expected incomplete RTO input to fail")

    copy_raw_reports.assert_not_called()
    assert not (tmp_path / "VAHAN_KERALA_MASTER1.xlsm").exists()


def test_validate_generated_dashboard_requires_all_rtos(tmp_path):
    for name in [
        "VAHAN_KERALA_MASTER1.xlsm",
        "MASTER_DATA_2026.xlsm",
        "RTO_2025_2026_Combined.csv",
    ]:
        (tmp_path / name).write_bytes(b"test")

    (tmp_path / "index.html").write_text("const rawRecords = [];", encoding="utf-8")

    (tmp_path / "rto_data.json").write_text(
        json.dumps({"Kollam": {"subRtoCodes": [f"KL{i}" for i in range(1, 86)]}}),
        encoding="utf-8",
    )

    try:
        integration.validate_generated_dashboard(tmp_path)
    except RuntimeError as exc:
        assert "contains 85 RTOs" in str(exc)
    else:
        raise AssertionError("Expected incomplete generated JSON to fail")


def test_integrate_preserves_existing_files_when_generation_fails(tmp_path):
    dashboard = tmp_path / "dashboard"
    reports = tmp_path / "reports" / "2026-09-14"
    dashboard.mkdir()
    reports.mkdir(parents=True)

    shutil.copy2(PROJECT_DIR / "RTO_MASTER.xlsm", dashboard / "RTO_MASTER.xlsm")
    shutil.copy2(PROJECT_DIR / "generate_rto_data.js", dashboard / "generate_rto_data.js")
    (dashboard / "MASTER_DATA_2026.xlsm").write_bytes(b"old-master")
    (dashboard / "RTO_2025_2026_Combined.csv").write_text(
        "Year,Year_Label,Month,RTO,RTO_clean,Mapped_Location,Maker,Cluster,Check_Eligible,Nos,Month_Num\n",
        encoding="utf-8",
    )
    (dashboard / "rto_data.json").write_text("old-json", encoding="utf-8")
    (dashboard / "index.html").write_text(
        "const rawRecords = []; const CURRENT_ASOF_DAY = 1;",
        encoding="utf-8",
    )

    for number in range(1, 87):
        make_xlsx_report(
            reports / f"KL{number}.xlsx",
            [["TOYOTA KIRLOSKAR MOTOR PVT LTD", 10, 20]],
        )

    original = {
        name: (dashboard / name).read_bytes()
        for name in ["MASTER_DATA_2026.xlsm", "RTO_2025_2026_Combined.csv", "rto_data.json", "index.html"]
    }

    with patch.object(integration, "run_dashboard_generator", side_effect=RuntimeError("generator failed")):
        try:
            integration.integrate(reports_dir=reports, dashboard_dir=dashboard)
        except RuntimeError as exc:
            assert "generator failed" in str(exc)
        else:
            raise AssertionError("Expected generator failure")

    assert {name: (dashboard / name).read_bytes() for name in original} == original
    assert not (dashboard / "VAHAN_KERALA_MASTER1.xlsm").exists()


def test_run_download_uses_automation_script_path():
    with patch("run_full_pipeline.subprocess.run") as run:
        run_full_pipeline.run_download()

    command = run.call_args.args[0]
    assert Path(command[-1]).name == "download_reports.py"
    assert Path(command[-1]).parent == AUTOMATION_DIR


def test_write_combined_csv_replaces_2026_and_preserves_2025(tmp_path):
    csv_path = tmp_path / "RTO_2025_2026_Combined.csv"
    csv_path.write_text(
        "Year,Year_Label,Month,RTO,RTO_clean,Mapped_Location,Maker,Cluster,Check_Eligible,Nos,Month_Num\n"
        "2025,2025,JAN,KL1,KL1,Kazhakkottam & Enchakkal,OLD MAKER,Cluster 3,Not Eligible,5,1\n"
        "2026,2026 (Jan-Aug),JAN,KL1,KL1,Kazhakkottam & Enchakkal,OLD 2026,Cluster 1,Required,99,1\n",
        encoding="utf-8",
    )
    rows = [{
        "RTO": "KL1",
        "NAME": "TRIVANDRUM RTO",
        "MAKER": "TOYOTA KIRLOSKAR MOTOR PVT LTD",
        "JAN": 10,
        "FEB": 20,
    }]

    integration.write_combined_csv(rows, csv_path, ["JAN", "FEB"])

    with csv_path.open(newline="", encoding="utf-8") as handle:
        output = list(csv.DictReader(handle))

    assert len(output) == 3
    assert output[0]["Year"] == "2025"
    assert {row["Year"] for row in output[1:]} == {"2026"}
    assert [row["Nos"] for row in output[1:]] == ["10", "20"]
    assert [row["Month_Num"] for row in output[1:]] == ["1", "2"]


def test_integration_updates_temp_dashboard_copy_and_runs_generator(tmp_path):
    dashboard = tmp_path / "performance-dashboard"
    reports = tmp_path / "reports" / "2026-09-14"
    dashboard.mkdir()
    reports.mkdir(parents=True)

    for name in [
        "RTO_MASTER.xlsm",
        "MASTER_DATA_2026.xlsm",
        "RTO_2025_2026_Combined.csv",
        "generate_rto_data.js",
        "package.json",
        "index.html",
    ]:
        shutil.copy2(PROJECT_DIR / name, dashboard / name)

    source_wb = openpyxl.load_workbook(dashboard / "MASTER_DATA_2026.xlsm", data_only=True, keep_vba=True)
    source_ws = source_wb[source_wb.sheetnames[0]]
    header = [cell.value for cell in source_ws[1]]
    by_rto = {}
    for row in source_ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            code = str(row[0]).replace(" ", "")
            by_rto.setdefault(code, []).append(row)

    for code, rows in by_rto.items():
        make_xlsx_report(reports / f"{code}.xlsx", [row[2:] for row in rows])

    integration.integrate(reports_dir=reports, dashboard_dir=dashboard, run_generator=True)

    assert (dashboard / "VAHAN_KERALA_MASTER1.xlsm").exists()
    assert (dashboard / "MASTER_DATA_2026.xlsm").exists()
    assert (dashboard / "RTO_2025_2026_Combined.csv").exists()
    assert (dashboard / "rto_data.json").exists()
    assert (dashboard / "vahan_uploads" / "2026-09-14").exists()

    html_records = integration.extract_index_records(dashboard / "index.html")
    assert any(r["Year"] == 2026 for r in html_records)

    manifest = json.loads((dashboard / "vahan_uploads" / "latest_manifest.json").read_text(encoding="utf-8"))
    assert manifest["report_file_count"] == 86
    assert manifest["rto_count"] == 86
    assert manifest["row_count"] == 1810

    data = json.loads((dashboard / "rto_data.json").read_text(encoding="utf-8"))
    assert "Kollam" in data
    assert data["Kollam"]["summary"]["tiv2026"] > 0


def test_validate_dashboard_snapshot_rejects_missing_rto_in_json(tmp_path):
    for name in [
        "MASTER_DATA_2026.xlsm",
        "RTO_MASTER.xlsm",
        "index.html",
    ]:
        shutil.copy2(PROJECT_DIR / name, tmp_path / name)

    shutil.copy2(PROJECT_DIR / "RTO_2025_2026_Combined.csv", tmp_path / "RTO_2025_2026_Combined.csv")
    (tmp_path / "rto_data.json").write_text(
        json.dumps({"Kollam": {"subRtoCodes": [f"KL{i}" for i in range(1, 86)]}}),
        encoding="utf-8",
    )

    try:
        validate_dashboard_snapshot(tmp_path)
    except RuntimeError as exc:
        assert "rto_data.json RTO coverage mismatch" in str(exc)
        assert "KL86" in str(exc)
    else:
        raise AssertionError("Expected incomplete JSON RTO coverage to fail")


def test_validate_dashboard_snapshot_accepts_current_committed_data():
    summary = validate_dashboard_snapshot(PROJECT_DIR)

    assert summary["rto_count"] == 86
    assert summary["cluster_count"] >= 1
