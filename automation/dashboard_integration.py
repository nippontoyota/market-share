import argparse
import csv
import json
import os
import re
import shutil
import subprocess
import tempfile
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import config


MONTHS = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]
MONTH_NUM = {month: index + 1 for index, month in enumerate(MONTHS)}

APPROVED_MAKERS = {
    "AUDI AG", "AUTOMOBILI LAMBORGHINI S.P.A", "BMW INDIA PVT LTD", "BYD INDIA PRIVATE LIMITED",
    "FORCE MOTORS LIMITED", "HINDUSTAN MOTORS LTD", "HONDA CARS INDIA LTD", "HYUNDAI MOTOR INDIA LTD",
    "ISUZU MOTORS INDIA PVT LTD", "JAGUAR LAND ROVER INDIA LIMITED", "JSW MG MOTOR INDIA PVT LTD",
    "KIA INDIA PRIVATE LIMITED", "MAHINDRA & MAHINDRA LIMITED", "MAHINDRA ELECTRIC AUTOMOBILE LTD",
    "MARUTI SUZUKI INDIA LTD", "MARUTI UDYOG LTD", "MERCEDES -BENZ AG", "MERCEDES-BENZ INDIA PVT LTD",
    "NISSAN MOTOR INDIA PVT LTD", "PORSCHE AG GERMANY", "RENAULT INDIA PVT LTD", "SKODA AUTO AS",
    "SKODA AUTO VOLKSWAGEN INDIA PVT LTD", "STELLANTIS AUTOMOBILES INDIA PVT LTD", "STELLANTIS INDIA PVT LTD",
    "TATA MOTORS PASSENGER VEHICLES LTD", "TATA PASSENGER ELECTRIC MOBILITY LTD",
    "TESLA INDIA MOTORS AND ENERGY PVT LTD", "TOYOTA KIRLOSKAR MOTOR PVT LTD", "VINFAST AUTO INDIA PVT LTD",
    "VOLKSWAGEN AG", "VOLVO AUTO INDIA PVT LTD",
}

CLUSTER_MAPPING = {
    "KL1": "Kazhakkottam & Enchakkal", "KL16": "Kazhakkottam & Enchakkal", "KL19": "Kazhakkottam & Enchakkal",
    "KL20": "Kazhakkottam & Enchakkal", "KL21": "Kazhakkottam & Enchakkal", "KL22": "Kazhakkottam & Enchakkal",
    "KL74": "Kazhakkottam & Enchakkal", "KL81": "Kazhakkottam & Enchakkal",
    "KL2": "Kollam", "KL23": "Kollam", "KL24": "Kollam", "KL25": "Kollam", "KL80": "Kollam", "KL82": "Kollam",
    "KL3": "Pathanamthitta", "KL26": "Pathanamthitta", "KL62": "Pathanamthitta", "KL83": "Pathanamthitta",
    "KL4": "Kayamkulam", "KL29": "Kayamkulam", "KL31": "Kayamkulam", "KL61": "Kayamkulam", "KL66": "Kayamkulam",
    "KL5": "Kottayam", "KL33": "Kottayam", "KL34": "Kottayam", "KL35": "Kottayam", "KL36": "Kottayam", "KL67": "Kottayam",
    "KL6": "Muvattupuzha", "KL17": "Muvattupuzha", "KL37": "Muvattupuzha", "KL38": "Muvattupuzha", "KL44": "Muvattupuzha",
    "KL68": "Muvattupuzha", "KL69": "Muvattupuzha",
    "KL7": "Kalamassery", "KL40": "Kalamassery", "KL41": "Kalamassery", "KL42": "Kalamassery", "KL63": "Kalamassery",
    "KL8": "Trichur", "KL46": "Trichur", "KL48": "Trichur", "KL75": "Trichur",
    "KL32": "Nettoor", "KL39": "Nettoor", "KL43": "Nettoor",
    "KL45": "Irinjalakuda", "KL47": "Irinjalakuda", "KL64": "Irinjalakuda",
    "KL27": "Thiruvalla", "KL28": "Thiruvalla", "KL30": "Thiruvalla",
}


def normalize_rto_code(value):
    text = str(value or "").upper().strip()
    match = re.search(r"\bKL\s*0*(\d{1,2})\b", text)
    if match:
        return f"KL{int(match.group(1))}"
    return re.sub(r"\s+", "", text)


def clean_header(value):
    if isinstance(value, tuple):
        value = " ".join(str(part) for part in value if str(part) != "nan")
    return re.sub(r"\s+", " ", str(value or "")).strip().upper()


def clean_number(value):
    if value is None:
        return 0
    text = str(value).strip()
    if not text or text.lower() == "nan" or text in {"-", "--"}:
        return 0
    text = re.sub(r"[^0-9.-]", "", text)
    if not text:
        return 0
    return int(float(text))


def find_latest_reports_dir():
    reports_root = Path(config.REPORTS_DIR)
    if not reports_root.exists():
        raise FileNotFoundError(f"Reports folder not found: {reports_root}")

    dated_dirs = [path for path in reports_root.iterdir() if path.is_dir()]
    if not dated_dirs:
        raise FileNotFoundError(f"No dated report folders found under {reports_root}")

    return max(dated_dirs, key=lambda path: path.stat().st_mtime)


def read_name_map(dashboard_dir):
    import openpyxl

    workbook_path = Path(dashboard_dir) / "RTO_MASTER.xlsm"
    if not workbook_path.exists():
        return {}

    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True, keep_vba=True)
    ws = wb["MASTER_DATA"] if "MASTER_DATA" in wb.sheetnames else wb[wb.sheetnames[0]]
    names = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] and row[1]:
            names[normalize_rto_code(row[0])] = str(row[1]).strip()
    wb.close()
    return names


def read_report_frames(file_path):
    import pandas as pd

    suffix = file_path.suffix.lower()
    if suffix == ".csv":
        return [pd.read_csv(file_path)]

    frames = []
    try:
        excel = pd.read_excel(file_path, sheet_name=None, dtype=str)
        frames.extend(excel.values())
    except Exception:
        pass

    if frames:
        return frames

    try:
        return pd.read_html(file_path)
    except Exception as exc:
        raise ValueError(f"Could not read {file_path.name} as Excel or HTML table: {exc}") from exc


def candidate_tables(frame):
    tables = [frame]
    for row_index in range(min(5, len(frame.index))):
        promoted = frame.iloc[row_index + 1:].copy()
        promoted.columns = frame.iloc[row_index].tolist()
        tables.append(promoted)
    return tables


def extract_month_columns(columns):
    month_columns = {}
    for column in columns:
        header = clean_header(column)
        for month in MONTHS:
            if re.search(rf"\b{month}\b", header):
                month_columns[month] = column
    return month_columns


def row_value(row, index):
    return row[index] if index < len(row) else None


def extract_portal_layout_rows(frame, rto_code, rto_name):
    values = frame.fillna("").values.tolist()
    for month_row_index in range(min(10, len(values))):
        row = values[month_row_index]
        month_columns = {}
        for col_index, value in enumerate(row):
            header = clean_header(value)
            if header in MONTHS:
                month_columns[header] = col_index

        if len(month_columns) < 2:
            continue

        maker_column = None
        for heading_row in values[:month_row_index]:
            for col_index, value in enumerate(heading_row):
                if "MAKER" in clean_header(value) or "MANUFACTURER" in clean_header(value):
                    maker_column = col_index
                    break
            if maker_column is not None:
                break
        if maker_column is None:
            maker_column = max(min(month_columns.values()) - 1, 0)

        rows = []
        for source_row in values[month_row_index + 1:]:
            maker = str(row_value(source_row, maker_column) or "").strip()
            if not maker or maker.lower() == "nan":
                continue
            if "TOTAL" in maker.upper():
                continue

            month_values = {
                month: clean_number(row_value(source_row, column))
                for month, column in month_columns.items()
            }
            rows.append({
                "RTO": rto_code,
                "NAME": rto_name,
                "MAKER": maker.upper(),
                **month_values,
            })

        if rows:
            return rows

    return []


def extract_report_rows(file_path, rto_code, rto_name):
    rows = []
    for frame in read_report_frames(file_path):
        portal_rows = extract_portal_layout_rows(frame, rto_code, rto_name)
        if portal_rows:
            return portal_rows

        for table in candidate_tables(frame):
            columns = list(table.columns)
            maker_column = next((col for col in columns if "MAKER" in clean_header(col) or "MANUFACTURER" in clean_header(col)), None)
            month_columns = extract_month_columns(columns)
            if not maker_column or not month_columns:
                continue

            for _, source_row in table.iterrows():
                maker = str(source_row.get(maker_column, "")).strip()
                if not maker or maker.lower() == "nan":
                    continue
                if "TOTAL" in maker.upper():
                    continue

                month_values = {month: clean_number(source_row.get(column)) for month, column in month_columns.items()}

                rows.append({
                    "RTO": rto_code,
                    "NAME": rto_name,
                    "MAKER": maker.upper(),
                    **month_values,
                })

            if rows:
                return rows

    raise ValueError(f"No maker/month table found in {file_path.name}")


def load_reports(reports_dir, name_map):
    report_files = sorted(
        path for path in Path(reports_dir).iterdir()
        if path.is_file() and path.suffix.lower() in {".xls", ".xlsx", ".xlsm", ".csv", ".html", ".htm"}
    )
    if not report_files:
        raise FileNotFoundError(f"No report files found in {reports_dir}")

    all_rows = []
    used_report_files = []
    failures = []
    for file_path in report_files:
        rto_code = normalize_rto_code(file_path.stem)
        if not re.fullmatch(r"KL([1-9]|[1-7][0-9]|8[0-6])", rto_code):
            continue
        rto_name = name_map.get(rto_code, f"{rto_code} RTO")
        try:
            extracted_rows = extract_report_rows(file_path, rto_code, rto_name)
            all_rows.extend(row for row in extracted_rows if row["MAKER"] in APPROVED_MAKERS)
            used_report_files.append(file_path)
        except Exception as exc:
            failures.append(f"{file_path.name}: {exc}")

    if failures:
        raise RuntimeError("Some report files could not be parsed:\n" + "\n".join(failures))
    if not all_rows:
        raise RuntimeError("Report files were found, but no usable rows were extracted.")

    return all_rows, used_report_files


def aggregate_rows(rows):
    grouped = defaultdict(lambda: {month: 0 for month in MONTHS})
    names = {}
    for row in rows:
        key = (row["RTO"], row["MAKER"])
        names[row["RTO"]] = row["NAME"]
        for month in MONTHS:
            grouped[key][month] += int(row.get(month, 0) or 0)

    output = []
    for (rto, maker), month_values in sorted(grouped.items(), key=lambda item: (int(item[0][0][2:]), item[0][1])):
        output.append({
            "RTO": rto,
            "NAME": names.get(rto, f"{rto} RTO"),
            "MAKER": maker,
            **month_values,
        })
    return output


def active_months(rows):
    months = [month for month in MONTHS if any(int(row.get(month, 0) or 0) for row in rows)]
    return months or MONTHS[:datetime.now().month]


def backup_files(dashboard_dir, files_to_backup):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_dir = Path(dashboard_dir) / "backups" / "vahan_auto" / timestamp
    backup_dir.mkdir(parents=True, exist_ok=True)

    for relative_path in files_to_backup:
        source = Path(dashboard_dir) / relative_path
        if source.exists():
            destination = backup_dir / relative_path
            destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, destination)

    return backup_dir


def write_workbook(rows, output_path, sheet_name, months):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    headers = ["RTO", "NAME", "MAKER", *months]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, 0 if header in months else "") for header in headers])
    wb.save(output_path)


def read_existing_2025_rows(csv_path):
    if not csv_path.exists():
        return []
    with csv_path.open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        return [row for row in reader if row.get("Year") != "2026"]


def write_combined_csv(rows, csv_path, months):
    fieldnames = [
        "Year", "Year_Label", "Month", "RTO", "RTO_clean", "Mapped_Location",
        "Maker", "Cluster", "Check_Eligible", "Nos", "Month_Num"
    ]
    existing_rows = read_existing_2025_rows(csv_path)
    year_label = f"2026 (Jan-{months[-1].title()})"
    new_rows = []

    for row in rows:
        rto = row["RTO"]
        mapped_location = CLUSTER_MAPPING.get(rto, "Other / Unmapped")
        for month in months:
            new_rows.append({
                "Year": "2026",
                "Year_Label": year_label,
                "Month": month,
                "RTO": rto,
                "RTO_clean": rto,
                "Mapped_Location": mapped_location,
                "Maker": row["MAKER"],
                "Cluster": "Cluster 1",
                "Check_Eligible": "Required",
                "Nos": str(int(row.get(month, 0) or 0)),
                "Month_Num": str(MONTH_NUM[month]),
            })

    with csv_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(existing_rows)
        writer.writerows(new_rows)


def build_index_records(csv_path):
    records = []
    with Path(csv_path).open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            if row.get("Check_Eligible") != "Required":
                continue
            records.append({
                "Year": int(row["Year"]),
                "Year_Label": row["Year_Label"],
                "Month": row["Month"],
                "Mapped_Location": row["Mapped_Location"],
                "Maker": row["Maker"],
                "Nos": int(row["Nos"] or 0),
            })
    return records


def extract_index_records(html_path):
    html = Path(html_path).read_text(encoding="utf-8")
    match = re.search(r"const rawRecords = (\[.*?\]);", html, re.S)
    if not match:
        raise RuntimeError("Could not find the rawRecords dataset inside index.html.")
    return json.loads(match.group(1))


def _html_safe_json(records):
    text = json.dumps(records, ensure_ascii=False, separators=(",", ":"))
    return (
        text.replace("&", "\\u0026")
        .replace("<", "\\u003c")
        .replace(">", "\\u003e")
        .replace(" ", "\\u2028")
        .replace(" ", "\\u2029")
    )


def update_index_html(html_path, csv_path, as_of_day):
    """Refresh the dataset baked into index.html so the main dashboard page
    (which reads `rawRecords` client-side, not rto_data.json) reflects the
    latest Vahan pull instead of whatever day the page was last generated on.

    The "current, still-filling-up" month, its totals, and every derived
    label are computed client-side in index.html's own JS from `rawRecords`
    (see `currentPartialMonth`), so this only needs to update the raw dataset
    plus the one thing the data itself can't reveal: which day-of-month the
    Vahan pull was current through.
    """
    records = build_index_records(csv_path)
    if not records:
        raise RuntimeError("No 'Required' rows found while rebuilding index.html dataset.")

    html = Path(html_path).read_text(encoding="utf-8")

    new_blob = f"const rawRecords = {_html_safe_json(records)};"
    updated, count = re.subn(
        r"const rawRecords = \[.*?\];", lambda _match: new_blob, html, count=1, flags=re.S
    )
    if count != 1:
        raise RuntimeError("Could not locate the rawRecords dataset inside index.html.")

    updated, day_count = re.subn(
        r"const CURRENT_ASOF_DAY = \d+;",
        lambda _match: f"const CURRENT_ASOF_DAY = {int(as_of_day)};",
        updated,
        count=1,
    )
    if day_count != 1:
        raise RuntimeError("Could not locate CURRENT_ASOF_DAY inside index.html.")

    Path(html_path).write_text(updated, encoding="utf-8")


def copy_raw_reports(report_files, dashboard_dir, run_date):
    destination_dir = Path(dashboard_dir) / "vahan_uploads" / run_date
    destination_dir.mkdir(parents=True, exist_ok=True)
    for file_path in report_files:
        shutil.copy2(file_path, destination_dir / file_path.name)
    return destination_dir


def run_dashboard_generator(dashboard_dir):
    generator = Path(dashboard_dir) / "generate_rto_data.js"
    if not generator.exists():
        return
    subprocess.run(["node", str(generator)], cwd=dashboard_dir, check=True)


def validate_generated_dashboard(dashboard_dir, expected_rto_count=config.EXPECTED_RTO_COUNT):
    dashboard_dir = Path(dashboard_dir)
    required_files = [
        "VAHAN_KERALA_MASTER1.xlsm",
        "MASTER_DATA_2026.xlsm",
        "RTO_2025_2026_Combined.csv",
        "rto_data.json",
        "index.html",
    ]
    missing = [name for name in required_files if not (dashboard_dir / name).is_file()]
    if missing:
        raise RuntimeError(f"Generated dashboard files are missing: {', '.join(missing)}")

    csv_2026_total = sum(
        r["Nos"] for r in build_index_records(dashboard_dir / "RTO_2025_2026_Combined.csv") if r["Year"] == 2026
    )
    html_records = extract_index_records(dashboard_dir / "index.html")
    html_2026_total = sum(r["Nos"] for r in html_records if r["Year"] == 2026)
    if html_2026_total != csv_2026_total:
        raise RuntimeError(
            "index.html's embedded dataset is out of sync with RTO_2025_2026_Combined.csv "
            f"(index.html 2026 total={html_2026_total}, CSV 2026 total={csv_2026_total})."
        )

    try:
        data = json.loads((dashboard_dir / "rto_data.json").read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise RuntimeError(f"Generated rto_data.json is not valid JSON: {exc}") from exc

    if not isinstance(data, dict) or not data:
        raise RuntimeError("Generated rto_data.json does not contain dashboard clusters.")

    rto_codes = set()
    for cluster in data.values():
        if not isinstance(cluster, dict):
            raise RuntimeError("Generated rto_data.json contains an invalid cluster.")
        rto_codes.update(cluster.get("subRtoCodes", []))

    if len(rto_codes) != expected_rto_count:
        raise RuntimeError(
            f"Generated rto_data.json contains {len(rto_codes)} RTOs. "
            f"Expected {expected_rto_count}."
        )


def publish_staged_files(staging_dir, dashboard_dir, files_to_publish, backup_dir):
    dashboard_dir = Path(dashboard_dir)
    staging_dir = Path(staging_dir)
    published = []
    try:
        for relative_path in files_to_publish:
            source = staging_dir / relative_path
            destination = dashboard_dir / relative_path
            shutil.copy2(source, destination)
            published.append(relative_path)
    except Exception:
        for relative_path in reversed(published):
            backup = Path(backup_dir) / relative_path
            destination = dashboard_dir / relative_path
            if backup.exists():
                shutil.copy2(backup, destination)
            elif destination.exists():
                destination.unlink()
        raise


def write_manifest(dashboard_dir, reports_dir, report_files, rows, months, backup_dir):
    manifest = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "reports_dir": str(Path(reports_dir).resolve()),
        "report_file_count": len(report_files),
        "rto_count": len({row["RTO"] for row in rows}),
        "row_count": len(rows),
        "months": months,
        "backup_dir": str(backup_dir),
    }
    manifest_path = Path(dashboard_dir) / "vahan_uploads" / "latest_manifest.json"
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return manifest_path


def integrate(reports_dir=None, dashboard_dir=None, run_generator=True):
    reports_dir = Path(reports_dir) if reports_dir else find_latest_reports_dir()
    dashboard_dir = Path(dashboard_dir or config.PERFORMANCE_DASHBOARD_DIR)
    if not dashboard_dir.exists():
        raise FileNotFoundError(f"Performance dashboard folder not found: {dashboard_dir}")

    run_date = reports_dir.name
    print(f"Integrating reports from {reports_dir}")
    print(f"Target dashboard project: {dashboard_dir}")

    name_map = read_name_map(dashboard_dir)
    raw_rows, report_files = load_reports(reports_dir, name_map)
    rows = aggregate_rows(raw_rows)
    months = active_months(rows)
    rto_count = len({row["RTO"] for row in rows})

    if len(report_files) != config.EXPECTED_RTO_COUNT:
        raise RuntimeError(
            f"Parsed {len(report_files)} in-scope report files. "
            f"Expected {config.EXPECTED_RTO_COUNT}."
        )
    if rto_count != config.EXPECTED_RTO_COUNT:
        raise RuntimeError(
            f"Parsed {rto_count} RTOs from report data. "
            f"Expected {config.EXPECTED_RTO_COUNT}."
        )

    files_to_publish = [
        "VAHAN_KERALA_MASTER1.xlsm",
        "MASTER_DATA_2026.xlsm",
        "RTO_2025_2026_Combined.csv",
        "rto_data.json",
        "index.html",
    ]

    with tempfile.TemporaryDirectory(prefix="vahan_staging_") as temp_dir:
        staging_dir = Path(temp_dir)
        shutil.copy2(dashboard_dir / "RTO_MASTER.xlsm", staging_dir / "RTO_MASTER.xlsm")
        shutil.copy2(dashboard_dir / "generate_rto_data.js", staging_dir / "generate_rto_data.js")

        write_workbook(rows, staging_dir / "VAHAN_KERALA_MASTER1.xlsm", "RTO_DATA", months)
        write_workbook(rows, staging_dir / "MASTER_DATA_2026.xlsm", "MASTER_DATA_2026", months)
        shutil.copy2(dashboard_dir / "RTO_2025_2026_Combined.csv", staging_dir / "RTO_2025_2026_Combined.csv")
        write_combined_csv(rows, staging_dir / "RTO_2025_2026_Combined.csv", months)

        shutil.copy2(dashboard_dir / "index.html", staging_dir / "index.html")
        update_index_html(
            staging_dir / "index.html",
            staging_dir / "RTO_2025_2026_Combined.csv",
            as_of_day=datetime.now().day,
        )

        if run_generator:
            run_dashboard_generator(staging_dir)
        else:
            shutil.copy2(dashboard_dir / "rto_data.json", staging_dir / "rto_data.json")

        validate_generated_dashboard(staging_dir)

        backup_dir = backup_files(dashboard_dir, files_to_publish)
        print(f"Backed up existing dashboard files to {backup_dir}")

        copied_dir = copy_raw_reports(report_files, dashboard_dir, run_date)
        print(f"Copied raw reports to {copied_dir}")

        publish_staged_files(staging_dir, dashboard_dir, files_to_publish, backup_dir)
        manifest_path = write_manifest(dashboard_dir, reports_dir, report_files, rows, months, backup_dir)
        print("Regenerated and validated dashboard rto_data.json")

    print(f"Integration complete: {len(report_files)} files, {rto_count} RTOs, {len(rows)} maker rows.")
    print(f"Manifest written to {manifest_path}")


def main():
    parser = argparse.ArgumentParser(description="Integrate downloaded Vahan RTO reports into the performance dashboard.")
    parser.add_argument("--reports-dir", help="Dated reports folder. Defaults to the latest folder under reports/.")
    parser.add_argument("--dashboard-dir", help="Performance dashboard project folder.")
    parser.add_argument("--skip-generator", action="store_true", help="Update source files without running dashboard generator.")
    args = parser.parse_args()

    integrate(args.reports_dir, args.dashboard_dir, run_generator=not args.skip_generator)


if __name__ == "__main__":
    main()
